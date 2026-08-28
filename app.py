from __future__ import annotations

import base64
import csv
import io
import os
import re
from collections import OrderedDict
from dataclasses import dataclass
from datetime import date
from decimal import Decimal, InvalidOperation, ROUND_HALF_UP
from xml.sax.saxutils import escape

from flask import Flask, jsonify, render_template, request
from openpyxl import Workbook
from openpyxl.styles import Alignment, Border, Font, PatternFill, Side
from pypdf import PdfReader
from reportlab.lib import colors
from reportlab.lib.enums import TA_RIGHT
from reportlab.lib.pagesizes import A4
from reportlab.lib.styles import ParagraphStyle, getSampleStyleSheet
from reportlab.lib.units import mm
from reportlab.platypus import Paragraph, SimpleDocTemplate, Spacer, Table, TableStyle

app = Flask(__name__)
app.config["MAX_CONTENT_LENGTH"] = 25 * 1024 * 1024

MONEY = Decimal("0.01")
NUMBER = re.compile(r"^-?\d+(?:[.,]\d+)?$")
ORDER = re.compile(r"^Order\s+(.+?)\s+-\s+Ref\s*:\s*(.+?)\s*$", re.I)
INVOICE = re.compile(r"Invoice\s+#(\d+)", re.I)
SKIP_PREFIXES = (
    "Description Your ref.", "Delivery address", "TOTAL VAT", "Subtotal", "Shipping costs",
    "Invoice ", "Total quantity", "Current balance", "Your current", "Gold still", "Please ",
    "Bank coordinates", "Any delay", "Countermark", "Produits ", "RETENTION OF TITLE", "SAS with",
    "316 769", "Code APE", "Rue du", "Tél ", "VAT:", "Folio ", "Date ", "Customer ",
    "Exonération", "Daily gold", "Comireland", "Shamrock House", "Dundrum", "D14NW93", "Irlande",
    "Siren", "payment of", "remains vested", "Retrouvez", "(Our terms", "(1)", "(2)",
)


@dataclass
class Item:
    description: str
    reference: str
    size: str
    qty: int
    metal: Decimal
    mode: str
    labour: Decimal
    supplier_unit: Decimal


@dataclass
class OrderGroup:
    reference: str
    label: str
    items: list[Item]


def decimal_value(value: str) -> Decimal:
    return Decimal(value.replace(" ", "").replace(",", "."))


def parse_item(line: str) -> Item | None:
    tokens = line.split()
    mode_index = next((i for i in range(len(tokens) - 1, -1, -1) if tokens[i].lower() in {"a", "t", "c"}), -1)
    if mode_index < 3 or len(tokens) - mode_index not in {3, 4}:
        return None
    after = tokens[mode_index + 1:]
    if not all(NUMBER.match(value) for value in after):
        return None
    try:
        metal = decimal_value(tokens[mode_index - 1])
        qty = int(tokens[mode_index - 2])
    except (InvalidOperation, ValueError):
        return None
    ref_index = mode_index - 3
    size = ""
    if (ref_index >= 1 and "," in tokens[ref_index] and NUMBER.match(tokens[ref_index])
            and any(char.isdigit() for char in tokens[ref_index - 1])):
        size = tokens[ref_index]
        ref_index -= 1
    if ref_index < 0:
        return None
    # Empty PDF table cells disappear during text extraction. The rightmost
    # value before Size/Qty is Reference when present, otherwise Your Reference.
    reference = tokens[ref_index]
    description = " ".join(tokens[:ref_index]).strip()
    labour = decimal_value(after[0])
    supplier_unit = decimal_value(after[-2] if len(after) == 3 else after[0])
    return Item(description, reference, size, qty, metal, tokens[mode_index].lower(), labour, supplier_unit)


def extract_invoice(pdf_stream) -> tuple[str, OrderedDict[str, OrderGroup]]:
    reader = PdfReader(pdf_stream)
    text = "\n".join(page.extract_text() or "" for page in reader.pages)
    invoice_match = INVOICE.search(text)
    if not invoice_match:
        raise ValueError("The PDF does not contain a recognizable invoice number.")
    groups: OrderedDict[str, OrderGroup] = OrderedDict()
    current_ref: str | None = None
    for raw_line in text.splitlines():
        line = " ".join(raw_line.split()).strip()
        order_match = ORDER.match(line)
        if order_match:
            current_ref = order_match.group(2).strip().split()[0]
            groups.setdefault(current_ref, OrderGroup(current_ref, line, []))
            continue
        if (current_ref and line.lower().startswith("countermark")
                and groups[current_ref].items):
            groups[current_ref].items[-1].description += f"\n{line}"
            continue
        if not current_ref or not line or line.startswith(SKIP_PREFIXES):
            continue
        item = parse_item(line)
        if item:
            groups[current_ref].items.append(item)
    groups = OrderedDict((ref, group) for ref, group in groups.items() if group.items)
    if not groups:
        raise ValueError("No product rows grouped by Order Ref were found in this PDF.")
    return invoice_match.group(1), groups


def item_price(item: Item, gold_fix: Decimal, markup: Decimal, gold_markup: Decimal) -> Decimal:
    description = item.description.upper()
    is_18ct = "18CT" in description or "18K" in description
    is_9ct = "9CT" in description or "9K" in description
    if is_18ct or is_9ct:
        value = item.supplier_unit * (Decimal("1") + gold_markup / Decimal("100"))
        if item.mode.lower() == "t":
            gold_multiplier = Decimal("2") if is_18ct else Decimal("1")
            value += gold_fix * item.metal * gold_multiplier / item.qty
    else:
        value = item.supplier_unit * (Decimal("1") + markup / Decimal("100"))
    return value.quantize(MONEY, rounding=ROUND_HALF_UP)


def euro(value: Decimal) -> str:
    return f"€ {value:.2f}"


def csv_rows_for_order(invoice: str, group: OrderGroup, gold_fix: Decimal, markup: Decimal, gold_markup: Decimal) -> list[list]:
    priced = [(item, item_price(item, gold_fix, markup, gold_markup)) for item in group.items]
    grand_total = sum((unit * item.qty for item, unit in priced), Decimal("0")).quantize(MONEY)
    rows = [
        [],
        [f"Date {date.today():%d/%m/%Y}", "Gold Fix", euro(gold_fix), "", "", "", "", "", ""],
        [f"DELIVERY NOTE #{invoice}", "", "", "", "", "", "", "", ""],
        [group.label, "", "", "", "", "", "", "", ""],
        [],
        ["Description", "Reference", "Size", "Qty", "Metal", "Mode", "Unit", "Total"],
        ["", "", "", "", "", "", "", euro(grand_total)],
        [],
    ]
    for item, unit in priced:
        rows.append([
            item.description, item.reference, item.size, item.qty,
            f"{item.metal:.2f}".replace(".", ","), item.mode, euro(unit), euro(unit * item.qty),
        ])
    return rows


def csv_for_order(invoice: str, group: OrderGroup, gold_fix: Decimal, markup: Decimal, gold_markup: Decimal) -> bytes:
    output = io.StringIO(newline="")
    writer = csv.writer(output)
    writer.writerows(csv_rows_for_order(invoice, group, gold_fix, markup, gold_markup))
    return "\ufeff".encode("utf-8") + output.getvalue().encode("utf-8")


def pdf_for_order(invoice: str, group: OrderGroup, gold_fix: Decimal, markup: Decimal, gold_markup: Decimal) -> bytes:
    output = io.BytesIO()
    styles = getSampleStyleSheet()
    body = ParagraphStyle("InvoiceBody", parent=styles["BodyText"], fontName="Helvetica", fontSize=8, leading=10)
    small = ParagraphStyle("InvoiceSmall", parent=body, fontSize=7, leading=9)
    header = ParagraphStyle("InvoiceHeader", parent=small, textColor=colors.white, fontName="Helvetica-Bold")
    right = ParagraphStyle("InvoiceRight", parent=body, alignment=TA_RIGHT)

    def paragraph(value, style=body):
        text = escape(str(value)).replace("\n", "<br/>").replace("€", "&euro;")
        return Paragraph(text or "&#160;", style)

    def footer(canvas, document):
        canvas.saveState()
        canvas.setFont("Helvetica", 7)
        canvas.setFillColor(colors.HexColor("#657386"))
        canvas.drawString(15 * mm, 9 * mm, f"DELIVERY NOTE #{invoice} · {group.label}")
        canvas.drawRightString(A4[0] - 15 * mm, 9 * mm, f"Page {document.page}")
        canvas.restoreState()

    document = SimpleDocTemplate(
        output, pagesize=A4, leftMargin=15 * mm, rightMargin=15 * mm,
        topMargin=14 * mm, bottomMargin=15 * mm, title=f"DELIVERY NOTE #{invoice} - {group.reference}",
    )
    story = [
        Paragraph(f"DELIVERY NOTE #{escape(invoice)}", styles["Title"]),
        Spacer(1, 3 * mm),
        paragraph(f"Date {date.today():%d/%m/%Y}    ·    Gold Fix € {gold_fix:.2f}"),
        Spacer(1, 2 * mm),
        Paragraph(escape(group.label), styles["Heading2"]),
        Spacer(1, 4 * mm),
    ]
    table_rows = [[paragraph(value, header) for value in
                   ["Description", "Reference", "Size", "Qty", "Metal", "Mode", "Unit", "Total"]]]
    order_total = Decimal("0")
    for item in group.items:
        unit = item_price(item, gold_fix, markup, gold_markup)
        total = (unit * item.qty).quantize(MONEY)
        order_total += total
        table_rows.append([
            paragraph(item.description), paragraph(item.reference), paragraph(item.size, right),
            paragraph(item.qty, right), paragraph(f"{item.metal:.2f}", right), paragraph(item.mode),
            paragraph(euro(unit), right), paragraph(euro(total), right),
        ])
    table_rows.append([paragraph("Order total", right), "", "", "", "", "", "",
                       paragraph(euro(order_total), right)])
    table = Table(table_rows, colWidths=[63 * mm, 23 * mm, 12 * mm, 10 * mm, 14 * mm, 11 * mm, 22 * mm, 25 * mm],
                  repeatRows=1, hAlign="LEFT")
    table.setStyle(TableStyle([
        ("BACKGROUND", (0, 0), (-1, 0), colors.HexColor("#12233B")),
        ("TEXTCOLOR", (0, 0), (-1, 0), colors.white),
        ("FONTNAME", (0, 0), (-1, 0), "Helvetica-Bold"),
        ("GRID", (0, 0), (-1, -2), 0.35, colors.HexColor("#DCE3EB")),
        ("VALIGN", (0, 0), (-1, -1), "TOP"),
        ("BACKGROUND", (0, 1), (-1, -2), colors.white),
        ("BACKGROUND", (0, -1), (-1, -1), colors.HexColor("#EDF4FF")),
        ("SPAN", (0, -1), (6, -1)),
        ("BOX", (0, -1), (-1, -1), 0.5, colors.HexColor("#9FB0C4")),
        ("TOPPADDING", (0, 0), (-1, -1), 5),
        ("BOTTOMPADDING", (0, 0), (-1, -1), 5),
    ]))
    story.append(table)
    document.build(story, onFirstPage=footer, onLaterPages=footer)
    return output.getvalue()


def xlsx_for_orders(invoice: str, groups: OrderedDict[str, OrderGroup], gold_fix: Decimal,
                    markup: Decimal, gold_markup: Decimal) -> bytes:
    workbook = Workbook()
    used_sheet_names: set[str] = set()
    navy = "12233B"
    pale_blue = "EDF4FF"
    white = "FFFFFF"
    line = Side(style="thin", color="DCE3EB")
    currency_format = '€ #,##0.00'

    for index, (order_ref, group) in enumerate(groups.items()):
        base_name = re.sub(r"[\\/*?:\[\]]", "_", order_ref).strip() or f"Order {index + 1}"
        base_name = base_name[:31]
        sheet_name = base_name
        suffix = 2
        while sheet_name.casefold() in used_sheet_names:
            marker = f"_{suffix}"
            sheet_name = f"{base_name[:31 - len(marker)]}{marker}"
            suffix += 1
        used_sheet_names.add(sheet_name.casefold())

        sheet = workbook.active if index == 0 else workbook.create_sheet()
        sheet.title = sheet_name
        sheet.sheet_view.showGridLines = False
        sheet.freeze_panes = "A6"
        sheet.merge_cells("A1:H1")
        sheet["A1"] = f"DELIVERY NOTE #{invoice}"
        sheet["A1"].font = Font(name="Arial", size=18, bold=True, color=navy)
        sheet["A1"].alignment = Alignment(vertical="center")
        sheet.row_dimensions[1].height = 28
        sheet["A2"] = "Date"
        sheet["B2"] = date.today()
        sheet["B2"].number_format = "dd/mm/yyyy"
        sheet["D2"] = "Gold Fix"
        sheet["E2"] = float(gold_fix)
        sheet["E2"].number_format = currency_format
        sheet.merge_cells("A3:H3")
        sheet["A3"] = group.label
        sheet["A3"].font = Font(name="Arial", size=12, bold=True, color=navy)

        headers = ["Description", "Reference", "Size", "Qty", "Metal", "Mode", "Unit", "Total"]
        for column, value in enumerate(headers, 1):
            cell = sheet.cell(row=5, column=column, value=value)
            cell.font = Font(name="Arial", size=10, bold=True, color=white)
            cell.fill = PatternFill("solid", fgColor=navy)
            cell.alignment = Alignment(vertical="center")
        sheet.row_dimensions[5].height = 22

        first_item_row = 6
        for row, item in enumerate(group.items, first_item_row):
            unit = item_price(item, gold_fix, markup, gold_markup)
            values = [item.description, item.reference, item.size, item.qty, float(item.metal),
                      item.mode, float(unit), float(unit * item.qty)]
            for column, value in enumerate(values, 1):
                cell = sheet.cell(row=row, column=column, value=value)
                cell.font = Font(name="Arial", size=9)
                cell.border = Border(bottom=line)
                cell.alignment = Alignment(vertical="top", wrap_text=column == 1)
            sheet.cell(row=row, column=4).number_format = "0"
            sheet.cell(row=row, column=5).number_format = "0.00"
            sheet.cell(row=row, column=7).number_format = currency_format
            sheet.cell(row=row, column=8).number_format = currency_format
            if "\n" in item.description:
                sheet.row_dimensions[row].height = 30

        total_row = first_item_row + len(group.items)
        sheet.merge_cells(start_row=total_row, start_column=1, end_row=total_row, end_column=7)
        sheet.cell(row=total_row, column=1, value="Order total")
        sheet.cell(row=total_row, column=8, value=f"=SUM(H{first_item_row}:H{total_row - 1})")
        for column in range(1, 9):
            cell = sheet.cell(row=total_row, column=column)
            cell.fill = PatternFill("solid", fgColor=pale_blue)
            cell.font = Font(name="Arial", size=10, bold=True, color=navy)
            cell.border = Border(top=line, bottom=line)
        sheet.cell(row=total_row, column=1).alignment = Alignment(horizontal="right")
        sheet.cell(row=total_row, column=8).number_format = currency_format
        sheet.auto_filter.ref = f"A5:H{total_row - 1}"
        sheet.column_dimensions["A"].width = 48
        sheet.column_dimensions["B"].width = 18
        sheet.column_dimensions["C"].width = 11
        sheet.column_dimensions["D"].width = 8
        sheet.column_dimensions["E"].width = 11
        sheet.column_dimensions["F"].width = 9
        sheet.column_dimensions["G"].width = 14
        sheet.column_dimensions["H"].width = 14
        sheet.print_title_rows = "1:5"
        sheet.page_setup.orientation = "landscape"
        sheet.page_setup.fitToWidth = 1
        sheet.page_setup.fitToHeight = 0
        sheet.sheet_properties.pageSetUpPr.fitToPage = True
        sheet.print_area = f"A1:H{total_row}"

    output = io.BytesIO()
    workbook.save(output)
    return output.getvalue()


def safe_name(value: str) -> str:
    return re.sub(r"[^A-Za-z0-9._-]+", "_", value).strip("._") or "order"


@app.get("/")
def index():
    return render_template("index.html")


@app.post("/api/process")
def process_invoice():
    uploaded = request.files.get("invoice")
    if not uploaded or not uploaded.filename.lower().endswith(".pdf"):
        return jsonify(error="Choose a PDF invoice."), 400
    try:
        gold_fix = decimal_value(request.form.get("gold_fix", ""))
        markup = decimal_value(request.form.get("markup", "35"))
        gold_markup = decimal_value(request.form.get("gold_markup", "27"))
        if gold_fix <= 0 or markup < 0 or gold_markup < 0:
            raise InvalidOperation
    except (InvalidOperation, ValueError):
        return jsonify(error="Enter a valid positive gold fix and non-negative markup values."), 400
    try:
        invoice, groups = extract_invoice(uploaded.stream)
        files = []
        for order_ref, group in groups.items():
            content = csv_for_order(invoice, group, gold_fix, markup, gold_markup)
            pdf_content = pdf_for_order(invoice, group, gold_fix, markup, gold_markup)
            files.append({
                "order_ref": group.label,
                "filename": f"invoice_{invoice}_{safe_name(order_ref)}.csv",
                "content_base64": base64.b64encode(content).decode("ascii"),
                "pdf_filename": f"invoice_{invoice}_{safe_name(order_ref)}.pdf",
                "pdf_base64": base64.b64encode(pdf_content).decode("ascii"),
                "item_count": len(group.items),
                "preview_rows": csv_rows_for_order(invoice, group, gold_fix, markup, gold_markup),
            })
        xlsx_content = xlsx_for_orders(invoice, groups, gold_fix, markup, gold_markup)
        return jsonify(
            invoice=invoice,
            files=files,
            excel_filename=f"delivery_note_{invoice}_all_orders.xlsx",
            excel_base64=base64.b64encode(xlsx_content).decode("ascii"),
        )
    except (ValueError, InvalidOperation) as exc:
        return jsonify(error=str(exc)), 422
    except Exception:
        app.logger.exception("Invoice processing failed")
        return jsonify(error="The invoice could not be processed. Confirm it uses the supported company format."), 422


@app.get("/api/health")
def health():
    return jsonify(ok=True)


if __name__ == "__main__":
    app.run(host="0.0.0.0", port=int(os.getenv("PORT", "8000")), debug=os.getenv("FLASK_DEBUG") == "1")
